import io
import logging
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Style constants
HEADER_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E79")
SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color="2E75B6")
META_FONT = Font(name="Calibri", size=10, color="404040")
META_BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="404040")
TABLE_HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TABLE_FONT = Font(name="Calibri", size=10)
SUMMARY_FONT = Font(name="Calibri", size=10, bold=True)
FOOTER_FONT = Font(name="Calibri", size=9, italic=True, color="888888")

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
MATCHED_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
LOW_CONF_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
UNMATCHED_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
SUMMARY_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
BOTTOM_BORDER = Border(bottom=Side(style="medium", color="1F4E79"))

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


def generate_invoice_report(
    invoice: dict,
    items: list[dict],
    supplier: dict | None = None,
) -> bytes:
    """Generate a professional Excel report for an invoice.

    Args:
        invoice: Dict with invoice fields (invoice_number, invoice_date, etc.)
        items: List of dicts with item fields (extracted_name, matched, etc.)
        supplier: Optional dict with supplier fields (name, address, phone, etc.)

    Returns:
        Excel file content as bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Report"

    # Page setup
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    row = 1

    # --- Report Title ---
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    cell = ws.cell(row=row, column=1, value="INVOICE PROCESSING REPORT")
    cell.font = HEADER_FONT
    cell.alignment = LEFT
    row += 1

    # Separator line
    for col in range(1, 10):
        ws.cell(row=row, column=col).border = BOTTOM_BORDER
    row += 1

    # --- Invoice Metadata ---
    meta_pairs = [
        ("Invoice #:", invoice.get("invoice_number") or "N/A"),
        ("Supplier:", (supplier or {}).get("name") or invoice.get("supplier_name") or "Unknown"),
        (
            "Date:",
            _format_date(invoice.get("invoice_date")),
        ),
        ("Processed:", datetime.now(timezone.utc).strftime("%d %b %Y")),
        ("Payment Terms:", invoice.get("payment_terms") or "N/A"),
        ("Currency:", invoice.get("currency") or "MYR"),
    ]

    for label, value in meta_pairs:
        ws.cell(row=row, column=1, value=label).font = META_BOLD_FONT
        ws.cell(row=row, column=1).alignment = RIGHT
        c = ws.cell(row=row, column=2, value=value)
        c.font = META_FONT
        c.alignment = LEFT
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        row += 1

    if supplier:
        if supplier.get("address"):
            ws.cell(row=row, column=1, value="Address:").font = META_BOLD_FONT
            ws.cell(row=row, column=1).alignment = RIGHT
            c = ws.cell(row=row, column=2, value=supplier["address"])
            c.font = META_FONT
            c.alignment = WRAP
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            row += 1
        if supplier.get("phone"):
            ws.cell(row=row, column=1, value="Phone:").font = META_BOLD_FONT
            ws.cell(row=row, column=1).alignment = RIGHT
            c = ws.cell(row=row, column=2, value=supplier["phone"])
            c.font = META_FONT
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            row += 1

    row += 1  # blank row

    # --- Check if discounts exist ---
    has_discounts = any(
        item.get("extracted_discount") and float(item["extracted_discount"]) > 0
        for item in items
    )

    # --- Items Table Header ---
    headers = ["#", "Product Name", "Barcode", "Qty", "UOM", "Unit Price", "Total"]
    if has_discounts:
        headers.insert(6, "Discount")
    headers.append("Match")
    headers.append("Matched Product")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = TABLE_HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    row += 1

    # --- Items Data ---
    for item in items:
        col = 1

        # Line number
        ws.cell(row=row, column=col, value=item.get("line_number", "")).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = CENTER
        col += 1

        # Product name
        name = item.get("extracted_name") or ""
        ws.cell(row=row, column=col, value=name).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = WRAP
        col += 1

        # Barcode — prefer matched product barcode, fall back to extracted
        barcode = item.get("product_barcode") or item.get("extracted_barcode") or ""
        ws.cell(row=row, column=col, value=barcode).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = CENTER
        col += 1

        # Qty
        qty = _to_float(item.get("extracted_qty"))
        ws.cell(row=row, column=col, value=qty).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = CENTER
        ws.cell(row=row, column=col).number_format = "#,##0.##"
        col += 1

        # UOM
        ws.cell(row=row, column=col, value=item.get("extracted_uom") or "").font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = CENTER
        col += 1

        # Unit Price
        price = _to_float(item.get("extracted_unit_price"))
        ws.cell(row=row, column=col, value=price).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = RIGHT
        ws.cell(row=row, column=col).number_format = "#,##0.00"
        col += 1

        # Discount (if column exists)
        if has_discounts:
            disc = item.get("extracted_discount")
            if disc and float(disc) > 0:
                ws.cell(row=row, column=col, value=f"{float(disc)}%").font = TABLE_FONT
            else:
                ws.cell(row=row, column=col, value="-").font = TABLE_FONT
            ws.cell(row=row, column=col).alignment = CENTER
            col += 1

        # Total
        total = _to_float(item.get("extracted_total"))
        ws.cell(row=row, column=col, value=total).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = RIGHT
        ws.cell(row=row, column=col).number_format = "#,##0.00"
        col += 1

        # Match status
        matched = item.get("matched", False)
        confidence = _to_float(item.get("match_confidence"))
        method = item.get("match_method") or ""

        if matched and confidence and confidence >= 0.95:
            match_text = f"✓ {confidence:.2f}"
            row_fill = MATCHED_FILL
        elif matched or (confidence and confidence >= 0.75):
            match_text = f"~ {confidence:.2f}" if confidence else "~ ?"
            row_fill = LOW_CONF_FILL
        else:
            match_text = "✗"
            row_fill = UNMATCHED_FILL

        if method:
            match_text += f" ({method})"

        ws.cell(row=row, column=col, value=match_text).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = CENTER
        col += 1

        # Matched product name
        matched_product = item.get("product_name") or ""
        ws.cell(row=row, column=col, value=matched_product).font = TABLE_FONT
        ws.cell(row=row, column=col).alignment = WRAP

        # Apply row fill and borders
        for c in range(1, col + 1):
            ws.cell(row=row, column=c).fill = row_fill
            ws.cell(row=row, column=c).border = THIN_BORDER

        row += 1

    # --- Totals Section ---
    row += 1
    total_col = len(headers)  # last column
    currency = invoice.get("currency") or "MYR"

    if invoice.get("subtotal"):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_col - 1)
        ws.cell(row=row, column=1, value="Subtotal:").font = SUMMARY_FONT
        ws.cell(row=row, column=1).alignment = RIGHT
        c = ws.cell(row=row, column=total_col, value=f"{currency} {_to_float(invoice['subtotal']):,.2f}")
        c.font = SUMMARY_FONT
        c.alignment = RIGHT
        row += 1

    if invoice.get("discount_total") and float(invoice["discount_total"]) > 0:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_col - 1)
        ws.cell(row=row, column=1, value="Discount:").font = SUMMARY_FONT
        ws.cell(row=row, column=1).alignment = RIGHT
        c = ws.cell(row=row, column=total_col, value=f"{currency} {_to_float(invoice['discount_total']):,.2f}")
        c.font = SUMMARY_FONT
        c.alignment = RIGHT
        row += 1

    if invoice.get("grand_total"):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_col - 1)
        ws.cell(row=row, column=1, value="GRAND TOTAL:").font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        ws.cell(row=row, column=1).alignment = RIGHT
        c = ws.cell(row=row, column=total_col, value=f"{currency} {_to_float(invoice['grand_total']):,.2f}")
        c.font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        c.alignment = RIGHT
        for col in range(1, total_col + 1):
            ws.cell(row=row, column=col).border = Border(
                top=Side(style="double", color="1F4E79"),
                bottom=Side(style="double", color="1F4E79"),
            )
        row += 1

    # --- Summary Section ---
    row += 1
    matched_count = sum(1 for i in items if i.get("matched"))
    unmatched_count = len(items) - matched_count
    needs_review = sum(
        1 for i in items
        if not i.get("matched") and i.get("match_confidence") and float(i["match_confidence"]) >= 0.5
    )

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_col)
    summary_text = (
        f"Summary: {len(items)} items  |  "
        f"{matched_count} matched  |  "
        f"{unmatched_count} unmatched"
    )
    if needs_review:
        summary_text += f"  |  {needs_review} needs review"
    cell = ws.cell(row=row, column=1, value=summary_text)
    cell.font = SUMMARY_FONT
    cell.fill = SUMMARY_FILL
    cell.alignment = CENTER
    for col in range(1, total_col + 1):
        ws.cell(row=row, column=col).fill = SUMMARY_FILL
        ws.cell(row=row, column=col).border = THIN_BORDER
    row += 2

    # --- Column widths ---
    col_widths = {
        1: 5,    # #
        2: 40,   # Product Name
        3: 16,   # Barcode
        4: 8,    # Qty
        5: 10,   # UOM
        6: 12,   # Unit Price
        7: 12,   # Discount or Total
        8: 14,   # Total or Match
        9: 18,   # Match or Matched Product
        10: 35,  # Matched Product (if discount column)
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Write to bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    content = buffer.read()

    logger.info(
        "Generated Excel report: %d items, %d bytes",
        len(items),
        len(content),
    )

    return content


def _to_float(value) -> float | None:
    """Safely convert a value to float."""
    if value is None:
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def _format_date(value) -> str:
    """Format a date value for display."""
    if not value:
        return "N/A"
    if isinstance(value, str):
        try:
            from datetime import date
            d = date.fromisoformat(value)
            return d.strftime("%d %b %Y")
        except ValueError:
            return value
    if hasattr(value, "strftime"):
        return value.strftime("%d %b %Y")
    return str(value)
